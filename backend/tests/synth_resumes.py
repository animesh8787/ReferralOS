"""Synthetic resumes. Person, employers and links are invented (example.com / fictional names)."""
import io

HEADER = ["Alex Example", "alex.example@example.com | +1 415 555 0134 | github.com/alex-example | linkedin.com/in/alex-example"]
EDU = ["EDUCATION", "Example State University — B.Tech in Computer Science    Aug 2022 - May 2026",
       "Expected graduation May 2026. CGPA: 8.7/10"]
EXP = ["EXPERIENCE", "Backend Engineering Intern — Acme Test Co    Jun 2024 - Aug 2024",
       "- Built REST APIs in Python and FastAPI backed by PostgreSQL for an internal tools service.",
       "- Wrote pytest suites and added Docker based CI checks for the team."]
PROJ = ["PROJECTS", "Job Tracker — github.com/alex-example/job-tracker",
        "- Built a Flask and Redis service that tracks job applications.",
        "- Deployed with Docker on AWS."]
SKILLS = ["SKILLS", "Languages: Python, Java, SQL", "Cloud: AWS, Docker, k8s", "Tools: Git, Postgres, Rust"]
CERT = ["CERTIFICATIONS", "Example Cloud Practitioner (Sample Academy)"]
LINES = HEADER + EDU + EXP + PROJ + SKILLS + CERT


def _wrap(lines, n):
    import textwrap
    out = []
    for l in lines:
        out += textwrap.wrap(l, n, subsequent_indent="") or [""]
    return out


def pdf_single(lines=None) -> bytes:
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    y = 800
    for l in lines or LINES:
        c.setFont("Helvetica-Bold" if l.isupper() else "Helvetica", 10)
        c.drawString(40, y, l.replace("—", "-").replace("•", "-"))
        y -= 15
    c.save()
    return buf.getvalue()


def pdf_two_column() -> bytes:
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    y = 800
    for l in HEADER[:1]:
        c.setFont("Helvetica-Bold", 14)
        c.drawString(40, y, l)
        y -= 16
    c.setFont("Helvetica", 9)
    c.drawString(40, y, "alex.example@example.com | +1 415 555 0134")
    top = y - 30
    for x, block in ((40, EDU + EXP + CERT), (330, SKILLS + PROJ)):
        yy = top
        for l in _wrap([b.replace("—", "-") for b in block], 44):
            c.setFont("Helvetica-Bold" if l.isupper() else "Helvetica", 9)
            c.drawString(x, yy, l)
            yy -= 13 if not l.isupper() else 16
    c.save()
    return buf.getvalue()


def pdf_image_only() -> bytes:
    import fitz
    doc = fitz.open()
    page = doc.new_page()
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 200, 100), False)
    pix.clear_with(200)
    page.insert_image(fitz.Rect(50, 50, 250, 150), pixmap=pix)
    return doc.tobytes()


def docx_single() -> bytes:
    import docx
    d = docx.Document()
    for l in LINES:
        d.add_paragraph(l)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def docx_table_layout() -> bytes:
    import docx
    d = docx.Document()
    for l in HEADER:
        d.add_paragraph(l)
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "\n".join(EDU + EXP)
    t.rows[0].cells[1].text = "\n".join(SKILLS + PROJ)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()
