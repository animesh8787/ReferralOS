import copy
import io
import re
from datetime import date

import docx
import fitz
import pytest

from app.services import jd_analysis, matching, taxonomy
from app.services.tailoring import engine, render, score as rscore, validate

TODAY = date(2026, 6, 1)


def it(i, section, data, status="active", verified=False, evidence=None, flags=None, source="resume"):
    return {"id": i, "section": section, "data": data, "status": status, "verified": verified, "evidence_level": evidence, "flags": flags or [], "source": source}


def base_items(**o):
    items = [
        it("per", "personal", {"name": "Alex Example", "email": "alex.example@example.com", "phone": "+1 415 555 0134",
                               "links": ["github.com/alex-example", "linkedin.com/in/alex-example"]}),
        it("e1", "experience", {"title": "Backend Engineering Intern", "organization": "Acme Test Co", "start": "2024-06", "end": "2024-08",
                                "bullets": ["Built REST APIs in Python and FastAPI backed by PostgreSQL for an internal tools service.",
                                            "Cut report generation time by 30% by adding Redis caching.",
                                            "Wrote pytest suites and added Docker based CI checks for the team.",
                                            "Attended weekly team meetings and took notes for the group."]}),
        it("e2", "experience", {"title": "Teaching Assistant", "organization": "Example State University", "start": "2023-01", "end": "2023-05",
                                "bullets": ["Held office hours for an introductory course."]}),
        it("p1", "project", {"name": "Job Tracker", "stack": ["Flask", "Redis", "Docker"], "repo_url": "https://github.com/alex-example/job-tracker", "demo_url": None,
                             "description": ["Built a Flask and Redis service that tracks job applications.", "Deployed with Docker."], "outcomes": []}),
        it("p2", "project", {"name": "Vision Lab", "stack": ["PyTorch"], "repo_url": None, "demo_url": None, "description": ["Trained an image classifier with PyTorch."], "outcomes": []}),
        it("p3", "project", {"name": "Pottery Blog", "stack": ["HTML", "CSS"], "repo_url": None, "demo_url": None, "description": ["Static site about pottery."], "outcomes": []}),
        it("s1", "skill", {"name": "Python"}, evidence="USED_PROFESSIONALLY"), it("s2", "skill", {"name": "Java"}, evidence="CLAIMED_ONLY"),
        it("s3", "skill", {"name": "Kubernetes"}, evidence="CLAIMED_ONLY"), it("s4", "skill", {"name": "SQL"}, evidence="STUDIED"),
        it("d1", "education", {"institution": "Example State University", "degree": "B.Tech", "discipline": "Computer Science", "start": "2022-08", "end": "2026-05",
                               "expected": True, "gpa": "CGPA: 8.7/10", "coursework": ["Data Structures"]}),
        it("c1", "certification", {"name": "Example Cloud Practitioner"}),
    ]
    return items + o.get("extra", [])


JD_A = ("Requirements:\n- 2+ years of experience with Python\n- Experience with PostgreSQL\n- Strong Terraform experience required\n"
        "\nNice to have:\n- Kubernetes experience\n\nResponsibilities:\n- Build REST APIs and own services end to end\n")
JD_B = ("Requirements:\n- Experience with PyTorch and machine learning\n- Strong Python skills required\n\nResponsibilities:\n- Train and deploy models for production\n")
JOB = {"title": "Backend Engineer", "company": "Acme Test Co", "employment_type": "full_time", "domains": ["backend"]}


def build(items=None, jd=JD_A, job=None, options=None, llm=None):
    items = items if items is not None else base_items()
    a = jd_analysis.analyse(jd + "\n" + "x" * 120)
    m = matching.build_matrix(a, matching.summarize_profile(items, TODAY))
    return engine.build(items, a, m, job or JOB, options, TODAY, llm), a, m


def all_text(content):
    return " ".join(t for _, t in engine.render_nodes(content) if t) + " " + " ".join(g["group"] for g in content["skills"]) + " " + " ".join(
        s["text"] for s in (content["summary"] or {"sentences": []})["sentences"])


# ---------------------------------------------------------------- honesty
def test_unsupported_jd_technology_is_never_added():
    r, a, m = build()
    txt = all_text(r["content"]).lower()
    assert "terraform" not in txt  # required by the posting, absent from the profile
    names = {i["name"] for g in r["content"]["skills"] for i in g["items"]}
    assert "Terraform" not in names and {"Python", "Java", "Kubernetes", "SQL"} <= names  # only the user's own listed skills
    assert {x["term"]: x["status"] for x in m}["Terraform"] == "NO_EVIDENCE"
    assert "kubernetes" not in " ".join(s["text"] for s in r["content"]["summary"]["sentences"]).lower()  # listed-only skills are never featured


def test_no_invented_numbers_and_missing_metrics_become_questions():
    r, _, _ = build()
    facts = " ".join(validate.strings([i["data"] for i in base_items()]))
    nums = set(validate.NUM.findall(all_text(r["content"])))
    assert nums <= set(validate.NUM.findall(facts)), nums - set(validate.NUM.findall(facts))
    qs = [q for q in r["questions"] if q["type"] == "missing_metric"]
    assert qs and all("Do you have a number" in q["text"] for q in qs) and len(qs) <= 3
    assert not any("%" in b["text"] for e in r["content"]["projects"] for b in e["bullets"])


def test_contradictory_dates_block_the_affected_claims():
    items = base_items()
    items[2] = it("e2", "experience", {"title": "Teaching Assistant", "organization": "Example State University", "start": "2024-06", "end": "2023-01",
                                       "bullets": ["Taught Haskell to beginners."]}, flags=[{"type": "date_order", "note": "End before start"}])
    r, _, _ = build(items)
    txt = all_text(r["content"])
    assert "Teaching Assistant" not in txt and "Haskell" not in txt
    assert any(x["item_id"] == "e2" for x in r["notes"]["left_out"])
    rep = validate.validate_all(r["ledger"], items, {cid for cid, _ in engine.render_nodes(r["content"])})
    assert [c["item_id"] for c in rep["unresolved_contradictions"]] == ["e2"] and rep["ok_to_approve"]


def test_suggested_unconfirmed_items_never_reach_the_resume():
    items = base_items(extra=[it("g1", "project", {"name": "Secret GitHub Project", "stack": ["Rust"], "description": ["Written in Rust."]}, status="suggested", source="github"),
                              it("g2", "skill", {"name": "Haskell"}, status="suggested", source="github", evidence="USED_IN_PROJECT")])
    r, _, _ = build(items)
    txt = all_text(r["content"])
    assert "Secret GitHub Project" not in txt and "Haskell" not in txt and "Rust" not in txt


def test_ledger_is_complete_and_grounded():
    r, _, _ = build()
    ids = {c["id"] for c in r["ledger"]}
    nodes = engine.render_nodes(r["content"])
    assert nodes and all(cid in ids for cid, _ in nodes)
    rep = validate.validate_all(r["ledger"], base_items(), {cid for cid, _ in nodes})
    assert rep["ok_to_approve"] and rep["validated"] == rep["claims_total"] == len(nodes) and not rep["needs_confirmation"]
    assert all(c["fact_ids"] for c in r["ledger"])
    assert rep == validate.validate_all(copy.deepcopy(r["ledger"]), base_items(), {cid for cid, _ in nodes})  # idempotent


def test_master_profile_is_untouched_and_no_state_leaks_between_jobs():
    items = base_items()
    before = copy.deepcopy(items)
    a1, _, _ = build(items, JD_A)
    b, _, _ = build(items, JD_B, job={**JOB, "title": "ML Engineer", "domains": ["ml_ai"]})
    a2, _, _ = build(items, JD_A)
    assert items == before
    assert a1["content"] == a2["content"] and a1["ledger"] == a2["ledger"]  # deterministic, nothing carried over from job B
    assert "PyTorch" in all_text(b["content"]) and "PyTorch" not in " ".join(s["text"] for s in a1["content"]["summary"]["sentences"])
    assert b["content"]["summary"] != a1["content"]["summary"]


def test_summary_is_factual_short_and_free_of_superlatives():
    r, _, _ = build()
    s = " ".join(x["text"] for x in r["content"]["summary"]["sentences"])
    assert not re.search(r"expert|proven|track record|passionate|world-class|best|guru|ninja|\d+\+? years", s, re.I)
    assert "B.Tech student in Computer Science (expected May 2026)" in s and "Python" in s and "Terraform" not in s
    assert all(c["fact_ids"] for c in r["ledger"] if c["kind"] == "summary")
    none, _, _ = build([it("x", "skill", {"name": "Python"}, evidence="CLAIMED_ONLY")])
    assert none["content"]["summary"] is None  # not enough facts: no fluff


# ---------------------------------------------------------------- selection and ordering
def test_projects_ranked_by_relevance_and_stack_is_copied_not_extended():
    jd = "Requirements:\n- Experience with Docker and Redis\n- Experience with PyTorch\n"
    r, _, _ = build(jd=jd, options={"max_projects": 2})
    names = [p["name"] for p in r["content"]["projects"]]
    assert names == ["Job Tracker", "Vision Lab"] and "Pottery Blog" not in names  # the two that touch the role's skills
    assert [p["name"] for p in build(jd="Requirements:\n- Experience with Rust\n")[0]["content"]["projects"]] == ["Job Tracker", "Vision Lab"]  # nothing relevant: user's own order
    for p in r["content"]["projects"]:
        src = next(i for i in base_items() if i["id"] == p["item_id"])
        assert p["stack"] == src["data"]["stack"]
    assert [x for x in r["notes"]["left_out"] if x["section"] == "project"]


def test_skill_order_and_groups():
    r, _, _ = build()
    g = r["content"]["skills"]
    assert g[0]["items"][0]["name"] == "Python" or g[0]["group"] in ("Languages", "Databases")
    flat = {i["name"]: grp["group"] for grp in g for i in grp["items"]}
    assert flat["SQL"] == "Studied"  # studied skills are kept apart from used ones
    langs = next(grp for grp in g if grp["group"] == "Languages")["items"]
    assert [i["name"] for i in langs].index("Python") < [i["name"] for i in langs].index("Java")  # used before listed-only


def test_bullets_are_verbatim_and_most_relevant_kept_in_original_order():
    r, _, _ = build(options={"max_bullets": 2})
    e = r["content"]["experience"][0]
    src = base_items()[1]["data"]["bullets"]
    assert [b["text"] for b in e["bullets"]] == [src[k] for k in sorted(b["source_index"] for b in e["bullets"])]
    assert all(b["text"] in src for b in e["bullets"]) and e["omitted_bullets"] == 2
    assert src[3] not in [b["text"] for b in e["bullets"]]  # the irrelevant bullet is the one dropped


@pytest.mark.parametrize("job,items_fn,first_two", [
    ({**JOB, "employment_type": "internship"}, base_items, ["summary", "education"]),                                        # early-career: education up front
    (JOB, lambda: [i for i in base_items() if i["id"] != "d1"] + [it("e9", "experience", {"title": "Engineer", "organization": "Initech Sample", "start": "2020-01", "end": "present", "bullets": ["Built services in Python."]})], ["summary", "skills"]),
])
def test_section_order_depends_on_role_type(job, items_fn, first_two):
    r, _, _ = build(items_fn(), job=job)
    assert r["content"]["section_order"][:2] == first_two
    order = r["content"]["section_order"]
    if r["role_type"] == "experienced":
        assert order.index("experience") < order.index("projects") and order.index("education") > order.index("experience")


def test_ai_role_with_ml_projects_puts_projects_first():
    r, _, _ = build(jd=JD_B, job={**JOB, "title": "ML Engineer", "domains": ["ml_ai"], "employment_type": "full_time"}, items=[i for i in base_items() if i["id"] != "d1"])
    assert r["role_type"] == "ai_projects" and r["content"]["section_order"].index("projects") < r["content"]["section_order"].index("experience")


def test_education_is_copied_exactly():
    e = build()[0]["content"]["education"][0]
    assert (e["institution"], e["degree"], e["start"], e["end"], e["gpa"], e["expected"]) == ("Example State University", "B.Tech", "Aug 2022", "May 2026", "CGPA: 8.7/10", True)


# ---------------------------------------------------------------- LLM output is never trusted
class EvilLLM:
    name = "evil"

    def rewrite_bullet(self, bullet, jd_terms):
        if "REST APIs" in bullet:
            return bullet + " Deployed on Kubernetes serving 40% more traffic at Google."
        if "pytest" in bullet:
            return "Authored pytest suites and added Docker based CI checks for the whole team."
        return None


def test_llm_rewrites_are_validated_and_bad_ones_discarded():
    r, _, _ = build(llm=EvilLLM())
    texts = [b["text"] for e in r["content"]["experience"] for b in e["bullets"]]
    assert not any("Kubernetes" in t or "40%" in t or "Google" in t for t in texts)
    assert "Built REST APIs in Python and FastAPI backed by PostgreSQL for an internal tools service." in texts  # fell back to the original
    assert "Authored pytest suites and added Docker based CI checks for the whole team." in texts  # a harmless rewording passes
    rej = r["notes"]["rewrites_rejected"]
    assert len(rej) == 1 and any("technology" in w for w in rej[0]["why"]) and any("numbers" in w for w in rej[0]["why"]) and any("names" in w for w in rej[0]["why"])
    rep = validate.validate_all(r["ledger"], base_items(), {cid for cid, _ in engine.render_nodes(r["content"])})
    assert rep["ok_to_approve"]


def test_check_rewrite_rules():
    o = "Built REST APIs in Python."
    assert validate.check_rewrite(o, "Developed REST APIs in Python.") == []
    assert validate.check_rewrite(o, "Built REST APIs in Python and Go.")
    assert validate.check_rewrite(o, "Built 12 REST APIs in Python.")
    assert validate.check_rewrite(o, "Built REST APIs in Python " + "and more " * 40)


# ---------------------------------------------------------------- claim validation rules
def claim(text, ids, kind="bullet"):
    return {"id": "c1", "kind": kind, "text": text, "fact_ids": ids}


def test_validate_claim_rules():
    by = {i["id"]: i for i in base_items(extra=[it("g1", "project", {"name": "G", "stack": ["Rust"], "description": []}, status="suggested", source="github"),
                                                it("bad", "experience", {"title": "X", "organization": "Y", "bullets": ["Used Go."]}, flags=[{"type": "conflict", "note": ""}])])}
    assert validate.validate_claim(claim("Built REST APIs in Python.", ["e1"]), by)[0] == "validated"
    assert validate.validate_claim(claim("Built REST APIs in Python.", []), by)[0] == "excluded"
    assert validate.validate_claim(claim("Built REST APIs in Python.", ["nope"]), by)[0] == "excluded"
    assert validate.validate_claim(claim("Wrote Rust.", ["g1"]), by) == ("excluded", ["The supporting fact is only a suggestion you have not confirmed."])
    assert validate.validate_claim(claim("Used Go.", ["bad"]), by)[0] == "excluded"
    st, why = validate.validate_claim(claim("Built REST APIs in Python on Kubernetes, serving 5000 users.", ["e1"]), by)
    assert st == "needs_confirmation" and "Kubernetes" in why[0] and "5000" in why[1]
    assert validate.validate_claim(claim("Postgres", ["e1"], "skill"), by)[0] == "validated"
    assert validate.validate_claim(claim("Cut report time by 30%.", ["e1"]), by)[0] == "validated"  # number really is in the fact


# ---------------------------------------------------------------- resume match score
def test_resume_score_components_and_honest_suggestions():
    r, a, m = build()
    s1 = rscore.score(r["content"], a, m, r["questions"])
    s2 = rscore.score(copy.deepcopy(r["content"]), a, m, r["questions"])
    assert s1 == s2 and 0 <= s1["total"] <= 100 and set(s1["components"]) == set(rscore.COMPONENTS)
    assert abs(sum(c["points"] for c in s1["components"].values()) - s1["total"]) < 0.05 and "not an ATS score" in s1["label"]
    tips = {t["term"]: t for t in s1["improvements"] if t["term"]}
    assert tips["Terraform"]["type"] == "no_support" and "only if you genuinely" in tips["Terraform"]["action"] and tips["Terraform"]["est_gain"] == 0.0
    assert tips["Kubernetes"]["type"] == "confirm_skill" and tips["Kubernetes"]["est_gain"] > 0
    assert any(t["type"] == "add_number" for t in s1["improvements"])
    only = {k: 0 for k in rscore.COMPONENTS} | {"mandatory_coverage": 1}
    assert rscore.score(r["content"], a, m, [], {"weights": only})["total"] == rscore.score(r["content"], a, m, [], {"weights": only})["components"]["mandatory_coverage"]["points"]
    with pytest.raises(ValueError):
        rscore.normalize_weights({"technical_skills": 1})


def test_score_reflects_what_the_resume_shows():
    items = base_items()
    strong, a, m = build(items)
    bare = [i for i in items if i["id"] not in ("e1", "p1")]
    weak, a2, m2 = build(bare)
    assert rscore.score(strong["content"], a, m, [])["total"] > rscore.score(weak["content"], a2, m2, [])["total"]


# ---------------------------------------------------------------- rendering
def test_pdf_docx_txt_render_text_and_links():
    r, _, _ = build()
    c = r["content"]
    pdf, dx, txt = render.to_pdf(c), render.to_docx(c), render.to_txt(c)
    doc = fitz.open(stream=pdf, filetype="pdf")
    text = " ".join(" ".join(p.get_text().split()) for p in doc)
    assert doc.page_count == 1 and "Alex Example" in text and "Built REST APIs in Python" in text and "SKILLS" not in text and "Skills" in text
    uris = {l["uri"] for p in doc for l in p.get_links()}
    assert {"mailto:alex.example@example.com", "https://github.com/alex-example", "https://linkedin.com/in/alex-example", "https://github.com/alex-example/job-tracker"} <= uris
    d = docx.Document(io.BytesIO(dx))
    assert len(d.tables) == 0 and {x.target_ref for x in d.part.rels.values() if x.reltype.endswith("/hyperlink")} >= uris
    assert any(p.style.name == "Heading 2" and p.text == "Experience" for p in d.paragraphs) and any(p.style.name == "List Bullet" for p in d.paragraphs)
    assert txt.splitlines()[0] == "Alex Example" and "EXPERIENCE" in txt and "- Built REST APIs" in txt
    assert render.check_render(c, pdf, dx)["ok"]


def test_render_check_detects_lost_links_and_garbage():
    r, _, _ = build()
    c = r["content"]
    pdf, dx = render.to_pdf(c), render.to_docx(c)
    c2 = copy.deepcopy(c)
    c2["header"]["links"].append("example.org/not-in-the-files")
    bad = render.check_render(c2, pdf, dx)
    assert not bad["ok"] and any("lost links" in p for p in bad["problems"])
    assert not render.check_render(c, b"not a pdf", dx)["ok"] and not render.check_render(c, pdf, b"not a docx")["ok"]


def test_non_latin_text_does_not_crash_rendering():
    items = base_items()
    items[0]["data"]["name"] = "Zoë Łukasz"
    r, _, _ = build(items)
    pdf = render.to_pdf(r["content"])
    assert fitz.open(stream=pdf, filetype="pdf").page_count == 1
