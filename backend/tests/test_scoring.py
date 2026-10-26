import copy
import io
import random
import time
from datetime import date

import openpyxl
import pytest

from app.services import classify, scoring as sc
from app.services.scoring_run import from_jsonable
from tests import synth_contacts as sc_data
from tests.test_imports import do_import

TODAY = date(2026, 6, 1)


def inp(**o):
    base = {"position": "Senior Software Engineer", "company_name": "Acme Test Co", "company_key": "acme test", "role_category": "engineer",
            "role_keyword": "engineer", "seniority": "senior", "seniority_keyword": "senior", "domains": ["backend"],
            "connected_on": date(2025, 9, 1), "has_url": True, "has_email": False, "employment_verified": False, "open_roles": None, "today": TODAY,
            "user": {"target_company_keys": {"acme test"}, "excluded_company_keys": set(), "prior_employer_keys": set(),
                     "user_domains": ["backend", "ml_ai"], "has_targets": True}}
    u = o.pop("user", {})
    base["user"] = {**base["user"], **u}
    return {**base, **o}


# ---------------------------------------------------------------- classification
@pytest.mark.parametrize("title,cat,sen", [
    ("Senior Software Engineer", "engineer", "senior"), ("Software Engineer", "engineer", "mid"), ("Staff Engineer", "engineer", "staff_principal"),
    ("Software Engineering Intern", "engineer", "intern"), ("Junior Backend Developer", "engineer", "junior"),
    ("Engineering Manager", "eng_manager", "manager"), ("Technical Recruiter", "recruiter", "unknown"),
    ("Talent Acquisition Partner", "recruiter", "unknown"), ("Director of Engineering", "leadership", "director_plus"),
    ("VP Engineering", "leadership", "director_plus"), ("Product Manager", "product", "manager"), ("Senior Data Scientist", "engineer", "senior"),
    ("Barista", "other", "unknown"), ("", "unknown", "unknown"), (None, "unknown", "unknown"),
])
def test_role_and_seniority(title, cat, sen):
    c = classify.classify(title)
    assert (c["role_category"], c["seniority"]) == (cat, sen)


def test_domains():
    assert classify.domain_tags("Machine Learning Engineer, NLP") == {"ml_ai"}
    assert {"backend", "platform_cloud"} <= classify.domain_tags("Backend Platform Engineer")
    assert classify.domain_tags("") == set()


# ---------------------------------------------------------------- pure scoring
def test_scores_stay_in_range_and_are_deterministic():
    rnd = random.Random(7)
    for _ in range(300):
        i = inp(role_category=rnd.choice(list(sc.CAPABILITY) + ["unknown"]), seniority=rnd.choice(list(sc.SENIORITY) + ["unknown"]),
                connected_on=rnd.choice([None, date(2026, 1, 1), date(2018, 1, 1)]), company_key=rnd.choice(["", "acme test", "other"]),
                domains=rnd.sample(["backend", "ml_ai", "data"], rnd.randint(0, 3)), user={"has_targets": rnd.random() < 0.5})
        for pol in ("neutral", "exclude"):
            cfg = {**sc.DEFAULT_CONFIG, "missing_policy": pol}
            r = sc.score(i, cfg)
            assert 0 <= r["total"] <= 100 and 0 <= r["confidence"]["score"] <= 100
            assert r == sc.score(copy.deepcopy(i), cfg)
            assert abs(sum(c["points"] for c in r["components"].values()) - r["total"]) < 0.05


def test_perfect_inputs_reach_high_but_components_are_visible():
    r = sc.score(inp(role_category="eng_manager", seniority="senior", domains=["backend"], connected_on=date(2026, 3, 1)))
    assert r["total"] > 80 and {"points", "weight", "explanation", "value", "known"} <= set(r["components"]["company_relevance"])


def test_weight_changes_change_ranking():
    recruiter = inp(role_category="recruiter", position="Technical Recruiter", seniority="unknown", domains=[], company_key="other", company_name="Other")
    engineer = inp()
    base = lambda w: sc.score(recruiter, {**sc.DEFAULT_CONFIG, "weights": w})["total"] - sc.score(engineer, {**sc.DEFAULT_CONFIG, "weights": w})["total"]  # noqa: E731
    only_hiring = {**{k: 0 for k in sc.COMPONENTS}, "hiring_relevance": 1}
    only_company = {**{k: 0 for k in sc.COMPONENTS}, "company_relevance": 1}
    assert base(only_hiring) > 0 > base(only_company)  # recruiter wins on hiring, engineer wins on company fit


def test_weights_normalise_and_validate():
    w = sc.normalize_weights({k: 1 for k in sc.COMPONENTS})
    assert sum(w.values()) == pytest.approx(100, abs=1e-9)
    assert sc.normalize_weights(sc.DEFAULT_WEIGHTS) == pytest.approx(sc.DEFAULT_WEIGHTS)
    for bad in ({}, {**sc.DEFAULT_WEIGHTS, "company_relevance": -1}, {k: 0 for k in sc.COMPONENTS}, {**sc.DEFAULT_WEIGHTS, "extra": 5}):
        with pytest.raises(ValueError):
            sc.normalize_weights(bad)


def test_missing_data_policy_unknown_is_not_negative():
    known = sc.score(inp())
    unk = inp(user={"target_company_keys": set()})  # company relevance becomes unknown
    neutral = sc.score(unk, {**sc.DEFAULT_CONFIG, "missing_policy": "neutral"})
    excl = sc.score(unk, {**sc.DEFAULT_CONFIG, "missing_policy": "exclude"})
    cr = neutral["components"]["company_relevance"]
    assert cr["known"] is False and cr["points"] == 12.5 and "unknown" in cr["explanation"].lower()  # neutral: half of its 25 points
    assert excl["components"]["company_relevance"]["points"] == 0 and excl["components"]["company_relevance"]["effective_weight"] == 0
    assert sum(c["effective_weight"] for c in excl["components"].values()) == pytest.approx(100, abs=0.1)  # others renormalised
    assert neutral["confidence"]["score"] < known["confidence"]["score"] and excl["confidence"]["score"] < known["confidence"]["score"]
    not_target = sc.score(inp(user={"target_company_keys": {"zzz"}}))["components"]["company_relevance"]
    assert not_target["known"] and not_target["value"] == 0.2  # observed 'not a target' differs from unknown


def test_explanations_only_state_supported_facts():
    r = sc.score(inp(connected_on=None))["components"]
    assert "Connected" not in r["approachability"]["explanation"] and "past employer" not in r["approachability"]["explanation"]
    assert r["approachability"]["known"] is False
    r2 = sc.score(inp(user={"prior_employer_keys": {"acme test"}}))["components"]["approachability"]
    assert "past employer" in r2["explanation"] and r2["value"] == 1.0
    cap = sc.score(inp())["components"]["referral_capability"]["explanation"]
    assert "not confirmed referral access" in cap and "Senior Software Engineer" in cap
    assert "verified" not in sc.score(inp())["components"]["hiring_relevance"]["explanation"].lower()


def test_outdated_connections_are_labelled_inferred_and_reduce_potential():
    old, new = sc.score(inp(connected_on=date(2018, 1, 1))), sc.score(inp(connected_on=date(2026, 1, 1)))
    assert old["employment_status"] == "possibly_outdated" and new["employment_status"] == "likely_current"
    assert old["components"]["referral_capability"]["value"] < new["components"]["referral_capability"]["value"]
    assert "out of date" in old["components"]["referral_capability"]["explanation"]
    assert sc.score(inp(connected_on=None))["employment_status"] == "unknown"
    assert old["confidence"]["parts"]["verified_vs_inferred"] == 0.3  # we never claim verification


def test_excluded_company_scores_zero_and_is_not_tiered():
    r = sc.score(inp(user={"excluded_company_keys": {"acme test"}}))
    assert r["total"] == 0 and r["excluded"]


def test_tiers_have_no_duplicates_and_deterministic_tie_breaks():
    rnd = random.Random(1)
    rows = [{"id": f"{i:05d}", "name_norm": f"n{i % 50}", "connected_on": date(2025, 1, 1 + i % 28), "total": float(i % 20),
             "confidence": float(i % 7), "excluded": i % 97 == 0} for i in range(1300)]
    a = sc.assign_tiers(rows, [100, 400, 500])
    shuffled = rows[:]
    rnd.shuffle(shuffled)
    assert [(r["id"], r["rank"], r["tier"]) for r in a] == [(r["id"], r["rank"], r["tier"]) for r in sc.assign_tiers(shuffled, [100, 400, 500])]
    assert [r["rank"] for r in a] == list(range(1, 1301)) and len({r["id"] for r in a}) == 1300
    tiers = [r["tier"] for r in a if not r["excluded"]]
    assert tiers.count(1) == 100 and tiers.count(2) == 400 and tiers.count(3) == 500 and tiers.count(None) == 1300 - 14 - 1000
    assert all(r["tier"] is None for r in a if r["excluded"])
    best_first = [r for r in a if r["tier"] == 1]
    assert min(r["total"] for r in best_first) >= max(r["total"] for r in a if r["tier"] == 3)
    assert sc.assign_tiers(rows, [5])[4]["tier"] == 1 and sc.assign_tiers(rows, [5])[5]["tier"] is None


# ---------------------------------------------------------------- API
def prospects(a, q=""):
    return a.get(f"/api/prospects?page_size=200{q}").json()


def test_scoring_runs_after_import_and_is_reproducible(env, mk):
    a = mk("a@example.com")
    do_import(a, sc_data.linkedin_csv(40))
    p = prospects(a)
    assert p["total"] == 40 and p["unscored_contacts"] == 0
    assert [i["rank"] for i in p["items"]] == list(range(1, 41)) and all(i["tier"] == 1 for i in p["items"])
    assert all(p["items"][k]["score"] >= p["items"][k + 1]["score"] for k in range(39))
    top = a.get(f"/api/prospects/{p['items'][0]['contact_id']}").json()
    assert top["scoring_version"] == sc.SCORING_VERSION and top["components"]["company_relevance"]["explanation"]
    redo = sc.score(from_jsonable(top["inputs"]), top["config_snapshot"])
    assert redo["total"] == top["score"] and redo["confidence"]["score"] == top["confidence"]  # reproducible from stored inputs


def test_user_context_drives_scores_and_config_changes_apply_on_recompute(env, mk):
    a = mk("a@example.com")
    do_import(a, sc_data.linkedin_csv(36))
    before = {i["contact_id"]: i["score"] for i in prospects(a)["items"]}
    a.put("/api/profile/preferences", {"target_roles": ["Backend Engineer"], "target_companies": ["Northwind Labs"], "excluded_companies": ["Wonka Synthetic"]})
    assert a.post("/api/scoring/recompute", {}).status_code == 202
    after = prospects(a)["items"]
    nw = [i for i in after if i["company"].startswith("Northwind Labs")]
    assert nw and all(i["rank"] <= 3 * len(nw) + 5 for i in nw)
    wonka = [i for i in after if i["company"].startswith("Wonka Synthetic")]
    assert wonka and all(i["excluded"] and i["score"] == 0 and i["tier"] is None for i in wonka)
    assert any(abs(before[i["contact_id"]] - i["score"]) > 1 for i in after)
    r = a.put("/api/scoring/config", {"weights": {"company_relevance": 0, "referral_capability": 0, "role_relevance": 0, "seniority_context": 0, "approachability": 0, "hiring_relevance": 1}})
    assert r.json()["weights"]["hiring_relevance"] == 100
    a.post("/api/scoring/recompute", {})
    top = prospects(a)["items"][0]
    assert top["role_category"] in ("recruiter", "eng_manager")


@pytest.mark.parametrize("body", [{"weights": {"company_relevance": 1}}, {"weights": {k: -1 for k in sc.COMPONENTS}}, {"missing_policy": "yolo"}, {"tiers": []}, {"tiers": [0]}])
def test_bad_config_rejected(env, mk, body):
    assert mk("a@example.com").put("/api/scoring/config", body).status_code == 422


def test_filters_pagination_and_search(env, mk):
    a = mk("a@example.com")
    do_import(a, sc_data.linkedin_csv(60))
    assert prospects(a, "&role_category=recruiter")["total"] > 0 and all(i["role_category"] == "recruiter" for i in prospects(a, "&role_category=recruiter")["items"])
    assert prospects(a, "&tier=2")["total"] == 0 and prospects(a, "&tier=1")["total"] == 60
    assert prospects(a, "&min_score=999")["total"] == 0
    assert all("manager" in (i["position"] or "").lower() for i in prospects(a, "&q=manager")["items"])
    a.put("/api/scoring/config", {"tiers": [10, 20]})
    a.post("/api/scoring/recompute", {})
    assert [prospects(a, f"&tier={t}")["total"] for t in (1, 2)] == [10, 20]
    p = a.get("/api/prospects?page=2&page_size=25").json()
    assert [i["rank"] for i in p["items"]][0] == 26 and prospects(a, "&sort=name")["items"][0]["name"] <= prospects(a, "&sort=name")["items"][1]["name"]


def test_company_intelligence(env, mk):
    a = mk("a@example.com")
    do_import(a, sc_data.linkedin_csv(48))
    do_import(a, sc_data.to_csv(sc_data.HR_HEADER, [["Acme Test Co", "Hr Person", "Recruiter", "hr@acme-test.example.com", "", "", "", ""]], preamble=""), kind="hr", provenance="self_collected")
    cid = [c for c in a.get("/api/companies").json()["items"] if c["name"] == "Acme Test Co"][0]["id"]
    intel = a.get(f"/api/companies/{cid}/intelligence").json()
    assert intel["connections"] == 4 and intel["top_contacts"] and intel["hr_contacts"][0]["flags"] == ["role_account"]
    assert intel["verified_open_roles"]["known"] is False and len(intel["best_first_outreach"]) <= 3


def test_exports_csv_injection_and_xlsx_structure(env, mk):
    a = mk("a@example.com")
    rows = [["=HYPERLINK(\"http://x\")", "Evil", "https://www.linkedin.com/in/synthetic-1", "", "+cmd Corp", "@SUM(1)", "01 Mar 2023"]]
    do_import(a, sc_data.to_csv(sc_data.LI_HEADER, rows, preamble=""))
    csv_text = a.get("/api/exports/prospects.csv").text
    assert "'=HYPERLINK" in csv_text and "'+cmd Corp" in csv_text and "'@SUM" in csv_text and ",=HYPERLINK" not in csv_text
    wb = openpyxl.load_workbook(io.BytesIO(a.get("/api/exports/prospects.xlsx").content))
    ws = wb["Prospects"]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref and ws["A1"].font.bold and ws.max_row == 2
    assert ws["E2"].value.startswith("'=HYPERLINK")
    assert wb["Summary"]["A1"].value and any("Scoring version" == r[0].value for r in wb["Summary"].iter_rows())


def test_isolation_for_prospects_and_intelligence(env, mk):
    a, b = mk("a@example.com"), mk("b@example.com")
    do_import(a, sc_data.linkedin_csv(12))
    cid = a.get("/api/companies").json()["items"][0]["id"]
    pid = prospects(a)["items"][0]["contact_id"]
    assert prospects(b)["total"] == 0 and b.get(f"/api/prospects/{pid}").status_code == 404
    assert b.get(f"/api/companies/{cid}/intelligence").status_code == 404  # shared company row must not leak who knows it
    assert b.get("/api/companies").json()["total"] == 0 and b.get("/api/exports/prospects.csv").text.count("\n") == 1


@pytest.mark.slow
def test_score_11k(env, mk, capsys):
    a = mk("a@example.com")
    do_import(a, sc_data.linkedin_csv(11000))
    t = time.perf_counter()
    a.post("/api/scoring/recompute", {})
    s = time.perf_counter() - t
    t = time.perf_counter()
    r = a.get("/api/prospects?page=50&page_size=100&min_score=40")
    q = (time.perf_counter() - t) * 1000
    t = time.perf_counter()
    x = a.get("/api/exports/prospects.xlsx")
    xs = time.perf_counter() - t
    assert r.status_code == 200 and prospects(a)["total"] == 11000 and len(x.content) > 100000
    with capsys.disabled():
        print(f"\n[perf] score 11000 contacts: {s:.1f}s ({11000 / s:.0f}/s); prospects page query: {q:.0f}ms; xlsx export 11000 rows: {xs:.1f}s")
