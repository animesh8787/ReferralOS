import csv
import io
from datetime import datetime

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel
from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from .. import ratelimit, tasks
from ..audit import audit
from ..models.contacts import Company, Contact, EmailContact, SuppressionEntry
from ..models.core import User, UserSetting
from ..models.scoring import ContactScore
from ..security import current_user, scoped_db
from ..services import scoring as sc
from ..services.scoring_run import get_config
from .contacts import _names, page_args

router = APIRouter(prefix="/api", tags=["scoring"])


# ---------------------------------------------------------------- config
@router.get("/scoring/config")
def get_cfg(db: Session = Depends(scoped_db)):
    return {**get_config(db), "version": sc.SCORING_VERSION, "components": sc.COMPONENTS}


class CfgIn(BaseModel):
    weights: dict[str, float] | None = None
    missing_policy: str | None = None  # neutral | exclude
    tiers: list[int] | None = None
    stale_years: int | None = None


@router.put("/scoring/config")
def put_cfg(body: CfgIn, user: User = Depends(current_user), db: Session = Depends(scoped_db)):
    cur = get_config(db)
    if body.weights is not None:
        try:
            cur["weights"] = sc.normalize_weights(body.weights)
        except ValueError as e:
            raise HTTPException(422, str(e))
    if body.missing_policy is not None:
        if body.missing_policy not in ("neutral", "exclude"):
            raise HTTPException(422, "missing_policy must be neutral or exclude")
        cur["missing_policy"] = body.missing_policy
    if body.tiers is not None:
        if not body.tiers or any(t < 1 for t in body.tiers) or len(body.tiers) > 6:
            raise HTTPException(422, "tiers must be 1-6 positive sizes")
        cur["tiers"] = body.tiers
    if body.stale_years is not None:
        cur["stale_years"] = max(1, min(body.stale_years, 30))
    s = db.query(UserSetting).filter(UserSetting.key == "scoring.config").first()
    if s:
        s.value = cur
    else:
        db.add(UserSetting(key="scoring.config", value=cur))
    audit(db, user.id, "scoring_config_changed")
    return {**cur, "note": "Recalculate to apply to existing prospects."}


@router.post("/scoring/recompute", status_code=202)
def recompute(user: User = Depends(current_user), db: Session = Depends(scoped_db)):
    ratelimit.check(f"recompute:{user.id}", 30, 3600)
    t = tasks.enqueue(db, "score_contacts")
    return {"task_id": t.id, "status": t.status, "result": t.result}


# ---------------------------------------------------------------- prospects
def _row(s: ContactScore, c: Contact, names: dict, detail: bool = False) -> dict:
    out = {"contact_id": c.id, "rank": s.rank, "tier": s.tier, "score": s.total, "confidence": s.confidence, "excluded": s.excluded,
           "name": f"{c.first_name} {c.last_name}".strip(), "company": names.get(c.company_id) or c.company_raw, "company_id": c.company_id,
           "position": c.position, "role_category": c.role_category, "seniority": c.seniority, "linkedin_url": c.linkedin_url,
           "connected_on": c.connected_on, "employment_status": c.employment_status,
           "components": {k: {"points": v["points"], "weight": v["weight"], "known": v["known"]} for k, v in s.components.items()}}
    if detail:
        out |= {"components": s.components, "confidence_parts": s.confidence_parts, "scoring_version": s.scoring_version,
                "config_snapshot": s.config_snapshot, "inputs": s.inputs, "computed_at": s.computed_at}
    return out


def _filtered(db, q, tier, min_score, min_conf, role, seniority, company_id):
    base = db.query(ContactScore, Contact).join(Contact, Contact.id == ContactScore.contact_id)
    if tier:
        base = base.filter(ContactScore.tier == tier)
    if min_score is not None:
        base = base.filter(ContactScore.total >= min_score)
    if min_conf is not None:
        base = base.filter(ContactScore.confidence >= min_conf)
    if role:
        base = base.filter(Contact.role_category == role)
    if seniority:
        base = base.filter(Contact.seniority == seniority)
    if company_id:
        base = base.filter(Contact.company_id == company_id)
    if q:
        like = f"%{q.strip().lower()}%"
        base = base.filter(or_(Contact.name_norm.like(like), func.lower(Contact.company_raw).like(like), func.lower(Contact.position).like(like)))
    return base


@router.get("/prospects")
def prospects(q: str | None = None, tier: int | None = None, min_score: float | None = None, min_confidence: float | None = None,
              role_category: str | None = None, seniority: str | None = None, company_id: str | None = None,
              sort: str = "rank", page: int = 1, page_size: int = 50, db: Session = Depends(scoped_db)):
    base = _filtered(db, q, tier, min_score, min_confidence, role_category, seniority, company_id)
    order = {"rank": ContactScore.rank, "score": ContactScore.total.desc(), "confidence": ContactScore.confidence.desc(),
             "name": Contact.name_norm}.get(sort, ContactScore.rank)
    total = base.count()
    off, lim = page_args(page, page_size)
    rows = base.order_by(order, ContactScore.rank).offset(off).limit(lim).all()
    names = _names(db, [c.company_id for _, c in rows])
    stale = db.query(Contact.id).outerjoin(ContactScore, ContactScore.contact_id == Contact.id).filter(ContactScore.id.is_(None)).count()
    return {"total": total, "page": page, "page_size": lim, "unscored_contacts": stale, "items": [_row(s, c, names) for s, c in rows]}


@router.get("/prospects/{contact_id}")
def prospect(contact_id: str, db: Session = Depends(scoped_db)):
    r = db.query(ContactScore, Contact).join(Contact, Contact.id == ContactScore.contact_id).filter(Contact.id == contact_id).first()
    if not r:
        raise HTTPException(404, "Not found")
    return _row(r[0], r[1], _names(db, [r[1].company_id]), detail=True)


# ---------------------------------------------------------------- company intelligence
@router.get("/companies/{company_id}/intelligence")
def intelligence(company_id: str, db: Session = Depends(scoped_db)):
    co = db.query(Company).filter(Company.id == company_id).first()
    rows = db.query(ContactScore, Contact).join(Contact, Contact.id == ContactScore.contact_id).filter(Contact.company_id == company_id) \
        .order_by(ContactScore.rank).all()
    hr = db.query(EmailContact).filter(EmailContact.company_id == company_id).all()
    if not co or not (rows or hr or db.query(Contact.id).filter(Contact.company_id == company_id).first()):
        raise HTTPException(404, "Not found")  # a company only you know about via your own contacts
    names = {company_id: co.name}
    by = lambda *cats: [_row(s, c, names) for s, c in rows if c.role_category in cats]  # noqa: E731
    sup = {v for (v,) in db.query(SuppressionEntry.value).filter(SuppressionEntry.kind == "email")}
    first = [_row(s, c, names) for s, c in rows if c.role_category in ("engineer", "eng_manager", "recruiter") and not s.excluded]
    return {"company": {"id": co.id, "name": co.name}, "connections": db.query(Contact).filter(Contact.company_id == company_id).count(),
            "by_role": {k: sum(1 for _, c in rows if c.role_category == k) for k in {c.role_category for _, c in rows}},
            "relevant_engineers": by("engineer", "eng_manager")[:5], "recruiters": by("recruiter")[:5], "managers_and_leads": by("eng_manager", "leadership")[:5],
            "high_confidence_prospects": [_row(s, c, names) for s, c in rows if s.confidence >= 60][:5],
            "top_contacts": [_row(s, c, names) for s, c in rows][:5], "best_first_outreach": sorted(first, key=lambda r: -r["score"])[:3],
            "hr_contacts": [{"id": e.id, "name": e.name, "title": e.title, "email": e.email_norm, "flags": e.email_flags,
                             "suppressed": bool(e.email_norm and e.email_norm in sup)} for e in hr[:10]],
            "verified_open_roles": {"known": False, "note": "Job discovery has not been run for this company yet."}}


# ---------------------------------------------------------------- exports
def safe(v):
    """Neutralise spreadsheet formula injection from imported text."""
    s = "" if v is None else str(v)
    return "'" + s if s[:1] in ("=", "+", "-", "@", "\t", "\r") else s


HEAD = ["Rank", "Tier", "Score", "Confidence", "Name", "Company", "Position", "Role category", "Seniority", "Employment (inferred)",
        "Connected on", "LinkedIn URL"]


def _export_rows(db):
    rows = db.query(ContactScore, Contact).join(Contact, Contact.id == ContactScore.contact_id).order_by(ContactScore.rank).limit(50000).all()
    names = _names(db, [c.company_id for _, c in rows])
    return [[s.rank, s.tier or "", s.total, s.confidence, f"{c.first_name} {c.last_name}".strip(), names.get(c.company_id) or c.company_raw or "",
             c.position or "", c.role_category or "", c.seniority or "", c.employment_status, c.connected_on or "", c.linkedin_url or ""] for s, c in rows]


@router.get("/exports/prospects.csv")
def export_csv(user: User = Depends(current_user), db: Session = Depends(scoped_db)):
    ratelimit.check(f"export:{user.id}", 5, 3600)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(HEAD)
    for r in _export_rows(db):
        w.writerow([safe(x) if isinstance(x, str) else x for x in r])
    audit(db, user.id, "export_prospects", format="csv")
    return Response(buf.getvalue(), media_type="text/csv", headers={"Content-Disposition": 'attachment; filename="prospects.csv"'})


@router.get("/exports/prospects.xlsx")
def export_xlsx(user: User = Depends(current_user), db: Session = Depends(scoped_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    ratelimit.check(f"export:{user.id}", 5, 3600)
    cfg, rows = get_config(db), _export_rows(db)
    wb = Workbook()
    ws = wb.active
    ws.title = "Prospects"
    ws.append(HEAD)
    for r in rows:
        ws.append([safe(x) if isinstance(x, str) else x for x in r])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEAD))}{max(2, len(rows) + 1)}"
    for i, wd in enumerate([6, 6, 8, 11, 24, 26, 34, 14, 14, 20, 12, 44], start=1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    sm = wb.create_sheet("Summary")
    sm.append(["Referral prospects summary"])
    sm["A1"].font = Font(bold=True, size=14)
    sm.append(["Generated (UTC)", datetime.utcnow().strftime("%Y-%m-%d %H:%M")])
    sm.append(["Scoring version", sc.SCORING_VERSION])
    sm.append(["Missing-data policy", cfg["missing_policy"]])
    sm.append(["Weights", ", ".join(f"{k}={v}" for k, v in cfg["weights"].items())])
    sm.append(["Prospects", len(rows)])
    sm.append([])
    sm.append(["Tier", "Count"])
    for t in sorted({r[1] for r in rows if r[1] != ""}):
        sm.append([f"Tier {t}", sum(1 for r in rows if r[1] == t)])
    sm.append(["Untiered / excluded", sum(1 for r in rows if r[1] == "")])
    sm.append([])
    sm.append(["Note", "Score ranks outreach priority. Confidence says how well-supported it is. Neither is a probability of a referral."])
    sm.column_dimensions["A"].width = 24
    sm.column_dimensions["B"].width = 90
    out = io.BytesIO()
    wb.save(out)
    audit(db, user.id, "export_prospects", format="xlsx")
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="prospects.xlsx"'})
