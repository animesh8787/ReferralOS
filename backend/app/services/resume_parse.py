"""Resume text extraction (PDF/DOCX) and deterministic structuring.

Every extracted item keeps its source text span, and string fields are verified to occur in the raw text
(unsupported extractions are rejected and reported, never kept). No LLM is involved.
Limits (documented): no OCR (image-only PDFs are flagged), heuristics assume English, conventional headings.
"""
import io
import re
from dataclasses import dataclass, field

from . import taxonomy

MAX_BYTES = 5 * 1024 * 1024
BULLETS = "-•*·▪●◦–"


class ParseError(Exception):
    pass


# ---------------------------------------------------------------- text extraction
def extract_text(data: bytes, filename: str) -> tuple[str, dict]:
    """-> (text, info). info: parser, pages, columns, image_only, warnings."""
    if not data:
        raise ParseError("The file is empty")
    if len(data) > MAX_BYTES:
        raise ParseError("The file is larger than 5 MB")
    name = filename.lower()
    if name.endswith(".pdf") or data[:5] == b"%PDF-":
        return _pdf(data)
    if name.endswith(".docx") or data[:2] == b"PK":
        return _docx(data)
    raise ParseError("Only PDF and DOCX files are supported")


def _pdf(data: bytes) -> tuple[str, dict]:
    import fitz  # PyMuPDF

    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as e:
        raise ParseError(f"Could not read the PDF: {type(e).__name__}") from e
    pages, columns = [], 0
    for page in doc:
        blocks = [b for b in page.get_text("blocks") if b[6] == 0 and b[4].strip()]
        text, two = _order_blocks(blocks, page.rect.width)
        columns = max(columns, 2 if two else 1)
        pages.append(text)
    text = "\n".join(pages).strip()
    info = {"parser": "pymupdf", "pages": len(doc), "columns": columns, "warnings": [], "image_only": False}
    if len(re.sub(r"\s", "", text)) < 20:
        info["image_only"] = True
        info["warnings"].append("No selectable text found. This looks like a scanned/image-only PDF; OCR is not "
                                "available, so please enter your details manually or upload a text-based file.")
        return text, info
    # header icons/words are hyperlinks: the visible text says only "GitHub", the URL lives in the annotation
    uris = list(dict.fromkeys(l["uri"] for page in doc for l in page.get_links() if l.get("uri")))
    if uris:
        text += "\nLINKS\n" + "\n".join(uris)
    return text, info


def _order_blocks(blocks, width: float) -> tuple[str, bool]:
    """Reading order; a two-column layout is detected when substantial blocks sit on both sides of the midline."""
    mid = width / 2
    left = [b for b in blocks if b[2] <= mid * 1.08]
    right = [b for b in blocks if b[0] >= mid * 0.92]
    spans = [b for b in blocks if b[0] < mid * 0.92 and b[2] > mid * 1.08]  # full-width (headers)
    two = len(left) >= 3 and len(right) >= 3 and len(spans) <= max(2, len(blocks) // 4)
    key = lambda b: (round(b[1], 0), b[0])  # noqa: E731
    if not two:
        return "\n".join(b[4].strip() for b in sorted(blocks, key=key)), False
    out = [b for b in sorted(spans, key=key) if b[1] < min(x[1] for x in left + right)]  # header above columns
    rest_spans = [b for b in spans if b not in out]
    for col in (sorted(left + [s for s in rest_spans], key=key), sorted(right, key=key)):
        out += col
    return "\n".join(b[4].strip() for b in out), True


def _docx(data: bytes) -> tuple[str, dict]:
    import docx

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise ParseError(f"Could not read the DOCX: {type(e).__name__}") from e
    lines = [p.text for p in d.paragraphs]
    for t in d.tables:  # multi-column resumes are often tables
        for row in t.rows:
            seen = set()
            for cell in row.cells:
                if id(cell._tc) not in seen:
                    seen.add(id(cell._tc))
                    lines.extend(p.text for p in cell.paragraphs)
    text = "\n".join(lines).strip()
    info = {"parser": "python-docx", "pages": None, "columns": 2 if d.tables else 1, "warnings": [], "image_only": False}
    if not text:
        raise ParseError("The document contains no text")
    return text, info


# ---------------------------------------------------------------- structuring
SECTIONS = {
    "education": ["education", "academic background", "academics"],
    "experience": ["experience", "work experience", "professional experience", "employment", "internships",
                   "internship experience", "work history", "research experience", "relevant experience",
                   "industry experience", "leadership experience", "teaching experience"],
    "projects": ["projects", "personal projects", "academic projects", "selected projects", "key projects"],
    "skills": ["skills", "technical skills", "skills & tools", "core competencies", "technologies", "tech stack"],
    "certifications": ["certifications", "certificates", "licenses", "licenses & certifications"],
    "awards": ["awards", "honors", "honours", "achievements", "honors & awards", "awards & achievements"],
    "publications": ["publications"],
    "summary": ["summary", "profile", "objective", "about me", "professional summary"],
    "links": ["links"],  # appended by _pdf from hyperlink annotations; read only by parse_personal
}
_HEAD = {a: k for k, v in SECTIONS.items() for a in v}
MONTH = r"(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|June?|July?|Aug(?:ust)?|Sep(?:t(?:ember)?)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)"
DATE = rf"(?:{MONTH}\.?,?\s+\d{{4}}|\d{{1,2}}/\d{{4}}|\d{{4}})"
RANGE = re.compile(rf"({DATE})\s*(?:-|–|—|to)\s*({DATE}|Present|Current|Now|Ongoing)", re.I)
EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
PHONE = re.compile(r"(?<!\w)\+?\d[\d\s().-]{8,}\d(?!\w)")
URL = re.compile(r"(?:https?://|www\.)[^\s,;)>\]]+|(?:linkedin\.com/in|github\.com)/[\w\-./]+", re.I)
TITLE_WORDS = ("engineer", "developer", "intern", "analyst", "manager", "lead", "scientist", "researcher", "architect",
               "consultant", "associate", "director", "designer", "administrator", "assistant", "specialist", "trainee",
               "programmer", "sde", "founder")
ORG_WORDS = ("university", "college", "institute", "school", "academy")
DEGREE = re.compile(r"\b(b\.?\s?tech|b\.?\s?e\.?|b\.?\s?s\.?c?|bachelor[^,\n]*|m\.?\s?tech|m\.?\s?s\.?c?|master[^,\n]*|ph\.?d\.?|"
                    r"doctorate|diploma|mba|b\.?\s?a\.?|m\.?\s?a\.?)\b", re.I)
GPA = re.compile(r"(?:c?gpa)\s*[:\-]?\s*(\d+(?:\.\d+)?)(?:\s*/\s*(\d+(?:\.\d+)?))?", re.I)
MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


@dataclass
class Item:
    section: str
    data: dict
    span: tuple[int, int]
    confidence: float = 0.7
    flags: list = field(default_factory=list)
    evidence_level: str | None = None


def iso_date(s: str) -> str | None:
    """'Jun 2024' -> '2024-06', '06/2024' -> '2024-06', '2024' -> '2024', 'Present' -> 'present'."""
    s = s.strip().rstrip(",")
    if s.lower() in ("present", "current", "now", "ongoing"):
        return "present"
    m = re.match(rf"({MONTH})\.?,?\s+(\d{{4}})", s, re.I)
    if m:
        return f"{m.group(2)}-{MONTHS[m.group(1)[:3].lower()]:02d}"
    m = re.match(r"(\d{1,2})/(\d{4})", s)
    if m:
        return f"{m.group(2)}-{int(m.group(1)):02d}"
    return s if re.fullmatch(r"\d{4}", s) else None


def _lines(text: str):
    pos = 0
    for raw in text.splitlines(keepends=True):
        yield raw.rstrip("\r\n"), pos
        pos += len(raw)


def split_sections(text: str):
    head, sections, cur = [], {}, None
    for line, off in _lines(text):
        key = re.sub(r"[^a-z& ]", "", line.strip().lower()).strip()
        if key in _HEAD and len(line.strip()) < 40:
            cur = _HEAD[key]
            sections.setdefault(cur, [])
        elif cur is None:
            head.append((line, off))
        elif line.strip():
            sections[cur].append((line, off))
    return head, sections


def _is_bullet(line: str) -> bool:
    s = line.lstrip()
    return bool(s) and s[0] in BULLETS and (len(s) == 1 or s[1] in " \t" or s[0] in "•·▪●◦")


def split_entries(lines):
    """Group lines into entries: header lines + bullets. Wrapped bullet text is joined onto the previous bullet."""
    entries, cur = [], None
    for line, off in lines:
        s = line.strip()
        if not s:
            continue
        if _is_bullet(line):
            if cur is None:
                cur = {"header": [], "bullets": [], "start": off, "end": off + len(line)}
                entries.append(cur)
            cur["bullets"].append(s.lstrip(BULLETS).strip())
            cur["end"] = off + len(line)
            continue
        wrapped = cur and cur["bullets"] and not RANGE.search(s) and (
            s[0].islower() or (len(cur["bullets"][-1]) >= 70 and not cur["bullets"][-1].endswith((".", "!", "?", ")"))))
        if wrapped:
            cur["bullets"][-1] += " " + s
            cur["end"] = off + len(line)
            continue
        new = cur is None or cur["bullets"] or (RANGE.search(s) and RANGE.search(" ".join(cur["header"])))
        if new:
            cur = {"header": [], "bullets": [], "start": off, "end": off}
            entries.append(cur)
        cur["header"].append(s)
        cur["end"] = off + len(line)
    return entries


def _split_header(header: str) -> list[str]:
    parts = re.split(r"\s+[—–|@]\s+|\s+-\s+|\s+at\s+|,\s+|\s{3,}|\t", header)
    return [p.strip(" ,-|") for p in parts if p.strip(" ,-|")]


def _parse_experience(e) -> Item | None:
    header = " ".join(e["header"])
    if not header:
        return None
    flags = []
    start = end = None
    m = RANGE.search(header)
    if m:
        start, end = iso_date(m.group(1)), iso_date(m.group(2))
        header = (header[:m.start()] + header[m.end():]).strip()
        if start and end and end != "present" and end < start:
            flags.append({"type": "date_order", "note": f"End date {end} is before start date {start}"})
    parts = _split_header(header)
    title = org = None
    for p in parts:
        if title is None and any(w in p.lower() for w in TITLE_WORDS):
            title = p
        elif org is None:
            org = p
    conf = 0.8 if title and org else 0.5
    if not title and parts:
        title = parts[0]
        org = org or (parts[1] if len(parts) > 1 else None)
        conf = 0.5
    data = {"title": title, "organization": org, "start": start, "end": end, "location": None,
            "bullets": e["bullets"], "metrics": []}
    return Item("experience", data, (e["start"], e["end"]), conf, flags)


def _parse_education(e) -> Item | None:
    blob = " ".join(e["header"] + e["bullets"])
    inst = next((p for p in _split_header(" ".join(e["header"])) if any(w in p.lower() for w in ORG_WORDS)), None)
    deg = DEGREE.search(blob)
    if not (inst or deg):
        return None
    gpa = GPA.search(blob)
    m = RANGE.search(blob)
    start = end = None
    if m:
        start, end = iso_date(m.group(1)), iso_date(m.group(2))
    else:
        years = re.findall(r"\b(?:19|20)\d{2}\b", blob)
        end = years[-1] if years else None
    data = {"institution": inst, "degree": deg.group(0).strip() if deg else None, "discipline": None,
            "start": start, "end": end, "expected": bool(re.search(r"expected|anticipated", blob, re.I)),
            "gpa": gpa.group(0).strip() if gpa else None, "coursework": []}
    cw = re.search(r"coursework[:\s]+(.+)", blob, re.I)
    if cw:
        data["coursework"] = [c.strip() for c in re.split(r"[,;]", cw.group(1)) if c.strip()]
    return Item("education", data, (e["start"], e["end"]), 0.75 if inst and deg else 0.5)


def _parse_project(e) -> Item | None:
    header = " ".join(e["header"])
    if not header:
        return None
    name = re.split(r"\s+[—–|]\s+|\s+-\s+|\s*\(", header)[0].strip(" :-")
    blob = " ".join(e["header"] + e["bullets"])
    urls = URL.findall(blob)
    stack = {c for c, _ in taxonomy.find_skills(blob)}
    for line in e["header"]:  # a tech-tag line ("FastAPI · React · SQLAlchemy"): keep the tags even when the taxonomy doesn't know them
        if "·" in line:
            stack |= {taxonomy.normalize(t) or t.strip() for t in line.split("·") if 0 < len(t.strip()) <= 30}
    stack = sorted(stack)
    data = {"name": name, "description": e["bullets"], "stack": stack,
            "repo_url": next((u for u in urls if "github.com" in u.lower()), None),
            "demo_url": next((u for u in urls if "github.com" not in u.lower()), None),
            "contribution": None, "outcomes": []}
    return Item("project", data, (e["start"], e["end"]), 0.7)


def _split_top(text: str) -> list[str]:
    """Split on , ; | and bullets, but not inside parentheses."""
    out, buf, depth = [], "", 0
    for ch in text:
        depth += (ch == "(") - (ch == ")" and depth > 0)
        if depth == 0 and ch in ",;|•·":
            out.append(buf)
            buf = ""
        else:
            buf += ch
    return out + [buf]


def _skill_groups(lines) -> list[tuple[str | None, str, int, int]]:
    """-> [(category, text, start, end)]. A short comma-less line followed by a list is a category label (PDF tables put it on its
    own line); a line that ends with ',' or '-', leaves a parenthesis open, or completes a known skill with the next word continues."""
    rows = [(l.strip().lstrip(BULLETS).strip(), o, o + len(l)) for l, o in lines if l.strip()]
    groups, k = [], 0
    while k < len(rows):
        t, a, b = rows[k]
        cat = None
        if ":" in t[:30]:
            cat, t = (x.strip() for x in t.split(":", 1))
        elif k + 1 < len(rows) and len(t) <= 35 and not re.search(r"[,;(]", t) and rows[k + 1][0].count(",") >= 2:
            cat, k = t, k + 1
            t, _, b = rows[k]
        while k + 1 < len(rows):
            nxt = rows[k + 1][0]
            last = _split_top(t)[-1].strip()
            first = nxt.split(",")[0].strip()
            if t.endswith("-"):
                t, k, b = t + nxt, k + 1, rows[k + 1][2]
            elif t.endswith(",") or t.count("(") > t.count(")") or taxonomy.normalize(f"{last} {first}"):
                t, k, b = f"{t} {nxt}", k + 1, rows[k + 1][2]
            else:
                break
        groups.append((cat, t, a, b))
        k += 1
    return groups


def _parse_skills(lines) -> list[Item]:
    items, seen = [], set()
    for cat, text, a, b in _skill_groups(lines):
        for tok in _split_top(text):
            tok = tok.strip().strip(".")
            names = [tok]
            m = re.fullmatch(r"(.+?)\s*\((.*)\)", tok)
            if m and "," in m.group(2):  # "AWS (Lambda, SQS)" is AWS plus each service
                names = [m.group(1)] + [x.strip() for x in m.group(2).split(",")]
            elif m and taxonomy.normalize(m.group(1)):  # "DynamoDB (single-table design)" is DynamoDB
                names = [m.group(1)]
            for n in names:
                n = n.strip().strip(".")
                canon = taxonomy.normalize(n)
                key = (canon or n).lower()
                if not n or len(n) > 40 or key in seen:
                    continue
                seen.add(key)
                items.append(Item("skill", {"name": canon or n, "category": taxonomy.category(canon) if canon else (cat or "other").lower(),
                                            "raw": n}, (a, b), 0.9 if canon else 0.6))
    return items


def _simple_lines(section: str, lines) -> list[Item]:
    items = []
    for e in split_entries(lines):
        names = [h for h in e["header"]] if not e["bullets"] else e["header"] + e["bullets"]
        items += [Item(section, {"name": n}, (e["start"], e["end"]), 0.7) for n in names if n.strip()]
    return items


def parse_personal(head, text: str, links_text: str = "") -> Item | None:
    data = {"name": None, "email": None, "phone": None, "location": None, "links": []}
    for line, _ in head:
        s = line.strip()
        if data["name"] is None and s and not EMAIL.search(s) and not URL.search(s) and not PHONE.search(s) \
                and re.fullmatch(r"[A-Za-z][A-Za-z.'\- ]{2,40}", s) and 2 <= len(s.split()) <= 4:
            data["name"] = s
    em = EMAIL.search(text[:600] + " " + links_text)
    data["email"] = em.group(0) if em else None
    ph = PHONE.search(text[:600])
    data["phone"] = ph.group(0).strip() if ph else None
    data["links"] = sorted(set(u.rstrip(".,") for u in URL.findall(text[:800] + " " + links_text)))
    if not any(data.values()):
        return None
    return Item("personal", data, (0, min(len(text), 600)), 0.7)


STRING_FIELDS = {"experience": ["title", "organization"], "education": ["institution"], "project": ["name"],
                 "personal": ["name", "email"], "certification": ["name"], "award": ["name"], "publication": ["name"]}


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", s).lower()


def traceable(item: Item, text: str) -> bool:
    hay = _norm(text)
    return all(not item.data.get(f) or _norm(str(item.data[f])) in hay for f in STRING_FIELDS.get(item.section, []))


def assign_evidence(items: list[Item]) -> None:
    """Skill evidence level from where the skill appears. Never auto-verified."""
    prof = set()
    proj = set()
    studied = set()
    for it in items:
        blob = " ".join(map(str, it.data.get("bullets", []))) + " " + str(it.data.get("title") or "")
        if it.section == "experience":
            prof |= taxonomy.skill_set(blob)
        elif it.section == "project":
            proj |= set(it.data.get("stack", []))
        elif it.section == "education":
            studied |= taxonomy.skill_set(" ".join(it.data.get("coursework", [])))
    ptext = " ".join(_norm(" ".join(map(str, it.data.get("bullets", []))) + " " + str(it.data.get("title") or "")) for it in items if it.section == "experience")
    jtext = " ".join(_norm(" ".join(map(str, it.data.get("description", []))) + " " + str(it.data.get("name") or "")) for it in items if it.section == "project")
    def implied(n, used):  # a used framework implies its language ("FastAPI" -> Python)
        return next((f for f in sorted(used) if n in taxonomy.IMPLIES.get(f, ())), None)

    for it in items:
        if it.section == "skill":
            n = it.data["name"]
            if not (n in prof or n in proj) and (f := implied(n, prof)):
                it.data["evidence_note"] = f"implied by {f} in your work experience"
                it.evidence_level = "USED_PROFESSIONALLY"
                continue
            if not (n in prof or n in proj) and (f := implied(n, proj)):
                it.data["evidence_note"] = f"implied by {f} in your projects"
                it.evidence_level = "USED_IN_PROJECT"
                continue
            word = re.compile(rf"(?<![\w]){re.escape(_norm(n))}(?![\w])")
            it.evidence_level = ("USED_PROFESSIONALLY" if n in prof or word.search(ptext) else "USED_IN_PROJECT" if n in proj or word.search(jtext)
                                 else "STUDIED" if n in studied else "CLAIMED_ONLY")


def flag_duplicates(items: list[Item]) -> None:
    seen = {}
    for it in items:
        if it.section == "experience":
            k = (_norm(it.data.get("title") or ""), _norm(it.data.get("organization") or ""))
            if k in seen:
                note = {"type": "duplicate", "note": "Same title and organization appears more than once"}
                it.flags.append(note)
                seen[k].flags.append(note)
            seen[k] = it


def structure(text: str) -> tuple[list[Item], dict]:
    """-> (items, report). report: rejected (untraceable), sections found."""
    head, sections = split_sections(text)
    items: list[Item] = []
    p = parse_personal(head, text, " ".join(l for l, _ in sections.get("links", [])))
    if p:
        items.append(p)
    for e in split_entries(sections.get("experience", [])):
        if it := _parse_experience(e):
            items.append(it)
    for e in split_entries(sections.get("education", [])):
        if it := _parse_education(e):
            items.append(it)
    for e in split_entries(sections.get("projects", [])):
        if it := _parse_project(e):
            items.append(it)
    items += _parse_skills(sections.get("skills", []))
    for sec, key in (("certifications", "certification"), ("awards", "award"), ("publications", "publication")):
        items += _simple_lines(key, sections.get(sec, []))
    if sections.get("summary"):
        items.append(Item("summary", {"text": " ".join(l.strip() for l, _ in sections["summary"])},
                          (sections["summary"][0][1], sections["summary"][-1][1]), 0.8))
    kept, rejected = [], []
    for it in items:
        (kept if traceable(it, text) else rejected).append(it)
    assign_evidence(kept)
    flag_duplicates(kept)
    return kept, {"rejected": [{"section": r.section, "data": r.data} for r in rejected],
                  "sections_found": sorted(sections), "item_count": len(kept)}
